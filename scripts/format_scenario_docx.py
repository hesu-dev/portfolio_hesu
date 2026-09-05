#!/usr/bin/env python3
"""Reflow a translated CoC scenario DOCX into session-friendly blocks.

The source file appears to be a PDF/OCR conversion where most page text is
collapsed into one paragraph. This script preserves the text, adds Roll20-style
chapter divider commands, and splits obvious markers into separate paragraphs.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ACCENT = "9F3535"

KNOWN_TITLE_PREFIXES = [
    "프롤로그",
    "2년 전",
    "골목 뒤쪽",
    "방아쇠에 닿아",
    "각성",
    "가을 비",
    "흰 꿈",
    "같은 날",
    "진료실",
    "형사 현장",
    "시부야역 앞",
    "검은 이형",
    "VS偽骸1",
    "거래",
    "상황 정리",
    "빨간 세계",
    "공동생활",
    "수사 재개",
    "가지 상황 청취",
    "상황 청취",
    "긴급 처치",
    "점심시간",
    "명의 환자 자료",
    "병원 습격",
    "가짜 해골 변이",
    "VS 가짜해골",
    "전투 종료",
    "레드 폐허",
    "만월의 밤",
    "보름달의 밤",
    "지명수배",
    "탈출병원",
    "키리시마 자택",
    "가짜 해골 개조",
    "그는 창문",
    "집단 자살",
    "가짜 시체 만남",
    "VS 가짜 시신",
    "키리시마 린도",
    "병원 탐색",
    "수술실",
    "자료실",
    "연구실",
    "신의 초대",
    "인류 구원",
    "정신 전이",
    "미래에서 과거로",
    "과거에서 미래로",
    "오다마키 코우조의 죽음",
    "VS 오다마키 코조",
    "수수께끼의 남자",
    "형사부 수사 1과",
    "제5 사건 현장",
    "조우",
    "古鳥軒",
    "동운대학 부속병원",
    "브리핑",
    "교로메",
    "수사본부 회의",
    "A-뱀하라 타츠미와 수사",
    "B-桃下冬香와 수사",
    "징후",
    "A-동승하는",
    "B-동승하지 않음",
    "VS 타카야마 하야토",
    "휴식",
    "의심",
    "A-2명 모두 휴식합니다",
    "불화",
    "VS 蛇原辰巳",
    "B-뱀하라 타츠미를 구속",
    "숨겨진 진실",
    "지옥의 문",
    "주요한 눈",
    "C-두쪽 모두 쏘지 않는다",
    "桃A-VS 桃下冬香",
    "B-VS 정의의 대행자",
    "풀가트리움의 밤",
]


def set_east_asia_font(run, font_name: str) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def set_style_font(style, *, font="Apple SD Gothic Neo", size=10.0, bold=False, color=None):
    style.font.name = font
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, size=10.2)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(3)

    for name in [
        "Scenario Macro",
        "Scenario Scene",
        "Scenario Branch",
        "Scenario Note",
        "Scenario Dialogue",
        "Scenario Handout",
    ]:
        if name not in styles:
            styles.add_style(name, 1)

    set_style_font(styles["Scenario Macro"], font="Menlo", size=7.8, color="FFFFFF")
    styles["Scenario Macro"].paragraph_format.space_before = Pt(10)
    styles["Scenario Macro"].paragraph_format.space_after = Pt(2)

    set_style_font(styles["Scenario Scene"], size=12.5, bold=True, color=ACCENT)
    styles["Scenario Scene"].paragraph_format.space_before = Pt(4)
    styles["Scenario Scene"].paragraph_format.space_after = Pt(5)

    set_style_font(styles["Scenario Branch"], size=10.2, bold=True, color="2F5D8C")
    styles["Scenario Branch"].paragraph_format.space_before = Pt(5)
    styles["Scenario Branch"].paragraph_format.space_after = Pt(3)

    set_style_font(styles["Scenario Note"], size=9.4, color="555555")
    styles["Scenario Note"].paragraph_format.left_indent = Inches(0.18)
    styles["Scenario Note"].paragraph_format.space_before = Pt(2)
    styles["Scenario Note"].paragraph_format.space_after = Pt(3)

    set_style_font(styles["Scenario Dialogue"], size=10.2, color="1E4F75")
    styles["Scenario Dialogue"].paragraph_format.left_indent = Inches(0.16)
    styles["Scenario Dialogue"].paragraph_format.space_before = Pt(3)
    styles["Scenario Dialogue"].paragraph_format.space_after = Pt(3)

    set_style_font(styles["Scenario Handout"], size=10.5, bold=True, color="333333")
    styles["Scenario Handout"].paragraph_format.space_before = Pt(7)
    styles["Scenario Handout"].paragraph_format.space_after = Pt(3)


def normalize_text(text: str) -> str:
    text = text.replace("\u3000", " ")
    text = text.replace("KP정보", "KP 정보")
    text = re.sub(r"대사\s*:", "대사:", text)
    text = re.sub(r"추가\s+대사\s*:", "추가 대사:", text)
    text = re.sub(r"\s+", " ", text).strip()

    # Split explicit scenario markers and speaker labels into copyable blocks.
    text = re.sub(r"(?<!^)\s*([▼◆※])", r"\n\1", text)
    text = re.sub(r"(?<!^)\s*(추가 대사:|대사:)", r"\n\1", text)
    text = re.sub(r"(?<!^)\s*(【[^】]{1,80}】)", r"\n\1", text)
    text = re.sub(r"(?<!^)\s*(●[^●\n]{1,40})", r"\n\1", text)
    text = re.sub(r"(?<!^)\s*(Result\b)", r"\n\1", text)
    text = re.sub(r"(?<!^)\s*(엔딩\s+[A-H](?:-\d+)?\s*[:：]?)", r"\n\1", text)

    # Common source artifact: the main body label and first scene are fused.
    text = re.sub(r"시나리오 본문\s+(00\s+프롤로그)", r"시나리오 본문\n\1", text)
    text = re.sub(r"\b(0\d-\d)(\d년\s+전)", r"\1 \2", text)
    text = re.sub(r"(수수께끼의 남자)(0?1)(?=\n?▼)", r"\n\2 \1", text)

    # Scene headings are often glued to the previous sentence during OCR/PDF
    # conversion: "...진행합시다. 02-3 진료실" or "...진행합시다. 레드 폐허 25".
    title_prefixes = "|".join(
        re.escape(prefix.replace("보름달의 밤", "만월의 밤"))
        for prefix in KNOWN_TITLE_PREFIXES
    )
    text = re.sub(
        rf"(?<=[.。!?！？」”』])\s*"
        rf"(?=(?:0\d|[1-4]\d)(?:-[A-Z0-9])?\s*(?:{title_prefixes}|20\d{{2}}년))",
        "\n",
        text,
    )
    text = re.sub(
        r"(?<=진행합시다\.)\s+(?=[가-힣A-Za-z][가-힣A-Za-z\s]{1,24}\s+\d{1,2}(?:\s|$))",
        "\n",
        text,
    )

    return text


def split_chunks(text: str) -> list[str]:
    lines: list[str] = []
    for raw in normalize_text(text).splitlines():
        line = raw.strip()
        if line:
            lines.append(line)
    return lines


def scene_number_value(scene_id: str) -> int | None:
    match = re.match(r"^(\d{1,2})", scene_id)
    if not match:
        return None
    return int(match.group(1))


def trim_title(title: str) -> str:
    title = re.sub(r"\s+", " ", title).strip(" .:：")
    title = title.replace("보름달의 밤", "만월의 밤")

    for prefix in KNOWN_TITLE_PREFIXES:
        if title.startswith(prefix):
            if prefix == "가지 상황 청취":
                return "상황 청취"
            if prefix == "명의 환자 자료":
                return "환자 자료"
            if prefix == "그는 창문":
                return "저녁의 각성"
            return prefix

    for stop in [" HO1", " HO2", " 2019년", " 2021년", " 대사:", " ▼", " ※"]:
        idx = title.find(stop)
        if idx > 0:
            title = title[:idx].strip()

    punct = re.search(r"[.。!?！？]", title)
    if punct and punct.start() <= 24:
        title = title[: punct.start() + 1].strip()

    if len(title) > 24:
        title = title[:24].rstrip() + "..."
    return title or "장면"


SCENE_PREFIX_RE = re.compile(
    r"^(?P<num>(?:0\d|[1-4]\d|\d)(?:-[A-Z0-9])?)\s*(?P<title>[0-9A-Za-z가-힣ぁ-んァ-ン一-龥々ー][^\n]{0,80})"
)


def parse_scene_prefix(line: str) -> tuple[str, str] | None:
    if line.startswith(("※", "▼", "◆")):
        return None
    match = SCENE_PREFIX_RE.match(line)
    if not match:
        return None

    scene_id = match.group("num")
    value = scene_number_value(scene_id)
    if value is None or value > 47:
        return None

    title_raw = match.group("title").strip()
    display_id = scene_id.zfill(2) if scene_id.isdigit() else scene_id
    letter_match = re.match(r"^(?:桃)?([A-Z])-", title_raw)
    title_for_display = title_raw
    if letter_match and "-" not in display_id:
        display_id = f"{display_id}{letter_match.group(1)}"
        title_for_display = re.sub(r"^(?:桃)?[A-Z]-", "", title_raw).strip()
    title = trim_title(title_for_display)

    # Avoid obvious false positives such as "23구", "11명의", "19시경" in
    # non-scene information blocks. Legitimate scene headings either use a
    # canonical title prefix, a zero-padded number, or a sub-scene suffix.
    has_known_title = any(
        title_raw.replace("보름달의 밤", "만월의 밤").startswith(prefix.replace("보름달의 밤", "만월의 밤"))
        for prefix in KNOWN_TITLE_PREFIXES
    )
    if not has_known_title and not scene_id.startswith("0") and "-" not in scene_id:
        return None

    return display_id, title


def parse_scene_suffix(line: str, next_line: str | None) -> tuple[str, str] | None:
    match = re.match(r"^(?P<title>[0-9A-Za-z가-힣ぁ-んァ-ン一-龥々ー][0-9A-Za-z가-힣ぁ-んァ-ン一-龥々ー\s-]{1,28})\s+(?P<num>\d{1,2})$", line)
    if not match:
        return None
    raw_title = match.group("title").strip()
    scene_id = match.group("num").zfill(2)
    has_known_title = any(raw_title.startswith(prefix) for prefix in KNOWN_TITLE_PREFIXES)
    if not next_line and not has_known_title:
        return None
    if next_line and not has_known_title and f"씬 {int(match.group('num'))}" not in next_line and f"씬 {scene_id}" not in next_line:
        return None
    return scene_id, trim_title(raw_title)


def parse_ending(line: str) -> tuple[str, str] | None:
    match = re.match(r"^(엔딩\s+[A-H](?:-\d+)?)(?:\s*[:：]\s*(.*))?", line)
    if not match:
        return None
    label = match.group(1).replace(" ", "")
    title = match.group(2) or match.group(1)
    return label, trim_title(title)


def macro_line(chapter: str, title: str) -> str:
    return (
        f'/desc [─────── CHAPTER {chapter} ───────](#" style="text-decoration:none; '
        f"color: #cbddee; background-color:#{ACCENT.lower()}; text-align:center; "
        f'display:block; padding:2px; box-shadow: 0px 8px 0px 15px #{ACCENT.lower()}; '
        f'font-size:12px; font-style: normal;)[{title}](#" style="text-decoration:none; '
        f"color: white; background-color:#{ACCENT.lower()}; text-align:center; display:block; "
        f'padding:2px; box-shadow:0px 8px 0px 15px #{ACCENT.lower()}; font-style: normal;)'
    )


def classify_style(line: str) -> str:
    if line.startswith(("/desc ", "【씬", "【엔딩")):
        return "Scenario Macro" if line.startswith("/desc ") else "Scenario Scene"
    if line.startswith("【HO") or line.startswith("【핸드아웃"):
        return "Scenario Handout"
    if line.startswith(("대사:", "추가 대사:")):
        return "Scenario Dialogue"
    if line.startswith("※") or line.startswith("▼KP 정보") or line.startswith("KP 정보"):
        return "Scenario Note"
    if line.startswith(("▼", "◆")):
        return "Scenario Branch"
    return "Normal"


def add_text_paragraph(doc: Document, text: str) -> None:
    style = classify_style(text)
    paragraph = doc.add_paragraph(style=style)
    if style == "Scenario Macro":
        set_paragraph_shading(paragraph, ACCENT)
    run = paragraph.add_run(text)
    if style == "Scenario Macro":
        set_east_asia_font(run, "Menlo")
    else:
        set_east_asia_font(run, "Apple SD Gothic Neo")


def add_scene_block(doc: Document, scene_id: str, title: str, emitted: set[str], *, ending=False) -> None:
    key = ("ending:" if ending else "scene:") + scene_id
    if key in emitted:
        return
    emitted.add(key)
    if ending:
        add_text_paragraph(doc, f"【{scene_id} : {title}】")
        add_text_paragraph(doc, macro_line(scene_id.replace("엔딩", "END-"), title))
    else:
        add_text_paragraph(doc, f"【씬{scene_id} : {title}】")
        add_text_paragraph(doc, macro_line(scene_id, title))


def collect_source_lines(source: Path) -> list[str]:
    doc = Document(source)
    lines: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if re.fullmatch(r"페이지\s+\d+", text):
            continue
        if paragraph.style.name == "Title":
            continue
        lines.extend(split_chunks(text))
    return lines


def build_docx(source: Path, output: Path) -> None:
    lines = collect_source_lines(source)

    doc = Document()
    configure_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("『데스트루도의 사자』 정리본")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor.from_string(ACCENT)
    set_east_asia_font(run, "Apple SD Gothic Neo")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run("챕터 구분 / 분기 / 대사 / 주석 기준 1차 정리")
    sub_run.font.size = Pt(10)
    sub_run.font.color.rgb = RGBColor.from_string("666666")
    set_east_asia_font(sub_run, "Apple SD Gothic Neo")

    emitted: set[str] = set()
    for idx, line in enumerate(lines):
        next_line = lines[idx + 1] if idx + 1 < len(lines) else None

        if line == "시나리오 본문":
            add_text_paragraph(doc, line)
            continue

        ending = parse_ending(line)
        if ending:
            add_scene_block(doc, ending[0], ending[1], emitted, ending=True)

        suffix_scene = parse_scene_suffix(line, next_line)
        if suffix_scene:
            add_scene_block(doc, suffix_scene[0], suffix_scene[1], emitted)
        else:
            prefix_scene = parse_scene_prefix(line)
            if prefix_scene:
                add_scene_block(doc, prefix_scene[0], prefix_scene[1], emitted)

        add_text_paragraph(doc, line)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build_docx(args.source, args.output)


if __name__ == "__main__":
    main()
