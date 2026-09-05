#!/usr/bin/env python3
"""Create copy-paste friendly Markdown/TXT scenario notes from the OCR DOCX."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from docx import Document


SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from format_scenario_docx import (  # noqa: E402
    classify_style,
    collect_source_lines,
    macro_line,
    parse_ending,
    parse_scene_prefix,
    parse_scene_suffix,
)


def extract_title(source: Path) -> str:
    doc = Document(source)
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text and paragraph.style.name == "Title":
            return text
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            return text
    return source.stem


def write_outputs(source: Path, markdown_output: Path, text_output: Path) -> None:
    lines = collect_source_lines(source)
    title = extract_title(source)
    markdown_output.parent.mkdir(parents=True, exist_ok=True)
    text_output.parent.mkdir(parents=True, exist_ok=True)

    emitted: set[str] = set()
    md_lines: list[str] = [
        f"# {title} 정리본",
        "",
        "챕터 구분 / 분기 / 대사 / 주석 기준 1차 정리",
        "",
    ]
    txt_lines: list[str] = [
        f"{title} 정리본",
        "챕터 구분 / 분기 / 대사 / 주석 기준 1차 정리",
        "",
    ]

    def emit_scene(label: str, title: str, *, ending: bool = False) -> None:
        key = ("ending:" if ending else "scene:") + label
        if key in emitted:
            return
        emitted.add(key)

        heading = f"【{label} : {title}】" if ending else f"【씬{label} : {title}】"
        md_lines.extend(["", f"## {heading}", "", "```text", macro_line(label, title), "```", ""])
        txt_lines.extend(["", heading, macro_line(label, title), ""])

    def emit_line(line: str) -> None:
        style = classify_style(line)
        if line.startswith("/desc "):
            md_lines.extend(["```text", line, "```", ""])
            txt_lines.extend([line, ""])
            return

        if style == "Scenario Handout":
            md_lines.extend(["", f"### {line}", ""])
            txt_lines.extend(["", line, ""])
        elif style == "Scenario Dialogue":
            md_lines.extend([f"> **{line}**", ""])
            txt_lines.extend([line, ""])
        elif style == "Scenario Note":
            md_lines.extend([f"> {line}", ""])
            txt_lines.extend([line, ""])
        elif style == "Scenario Branch":
            md_lines.extend(["", f"**{line}**", ""])
            txt_lines.extend(["", line, ""])
        else:
            md_lines.extend([line, ""])
            txt_lines.extend([line, ""])

    for idx, line in enumerate(lines):
        next_line = lines[idx + 1] if idx + 1 < len(lines) else None

        ending = parse_ending(line)
        if ending:
            label, title = ending
            emit_scene(label, title, ending=True)

        suffix_scene = parse_scene_suffix(line, next_line)
        if suffix_scene:
            scene_id, title = suffix_scene
            emit_scene(scene_id, title)
            continue
        else:
            prefix_scene = parse_scene_prefix(line)
            if prefix_scene:
                scene_id, title = prefix_scene
                emit_scene(scene_id, title)

        emit_line(line)

    markdown_output.write_text("\n".join(md_lines).strip() + "\n", encoding="utf-8")
    text_output.write_text("\n".join(txt_lines).strip() + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--md", required=True, type=Path)
    parser.add_argument("--txt", required=True, type=Path)
    args = parser.parse_args()
    write_outputs(args.source, args.md, args.txt)


if __name__ == "__main__":
    main()
